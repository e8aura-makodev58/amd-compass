"""Builds the ResponseCompass pitch deck as .pptx and a matching .pdf from one spec.

Backgrounds: slide 1 is the cover image (assets/AMDResponseCompass.png). Statement slides
use the wallpaper (assets/ResponseCompassWallpaper.png) with content on its empty left
half; content slides use a softened copy of it (generated into .cache/deck/) under white
cards, so the illustration shows only in the margins.

The diagrams use native PowerPoint shapes; every arrow is a connector glued to connection
sites on its two nodes (a:stCxn / a:endCxn), so moving a node in PowerPoint keeps the
arrow attached. The PDF is printed by Chromium from an HTML rendering of the same
geometry, using the app's vendored Inter font, so it matches the .pptx opened with Inter.

Run: PYTHONPATH=.cache/pylib LD_LIBRARY_PATH=<chromium libs> python3 docs/submission/build/build-deck.py
"""
import html
import subprocess
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parents[1]
ASSETS = ROOT / "assets"
CACHE = REPO / ".cache" / "deck"
OUT_PPTX = ROOT / "ResponseCompass-deck.pptx"
OUT_PDF = ROOT / "ResponseCompass-deck.pdf"
INTER = REPO / "frontend" / "vendor" / "fonts" / "inter-latin-wght-normal.woff2"

W, H = 13.333, 7.5
# The app's "Modern light" tokens (frontend/tokens.css), which the cover image also uses.
INK, SLATE, MUTED, INDIGO, INDIGO_DEEP, INDIGO_WASH, AMBER, AMBER_WASH, GREEN, LINE, BG, WHITE = (
    "111827", "475467", "667085", "4F46E5", "1E1B4B", "EEF0FF", "D97706", "FFF7ED", "16A34A", "D0D5DD", "F4F6FA", "FFFFFF")
FONT = "Inter"
COVER = ASSETS / "AMDResponseCompass.png"
WALLPAPER = ASSETS / "ResponseCompassWallpaper.png"
SOFT_WALLPAPER = CACHE / "wallpaper-soft.png"

# ---------------------------------------------------------------------------
# Slide spec. Coordinates are inches from the top-left corner. On wallpaper slides the
# illustration starts near x = 6.6, and the wallpaper's wordmark sits at x < 2.5, y > 6.8.
# ---------------------------------------------------------------------------

def text(x, y, w, h, value, size, bold=False, color=INK, align="left", spacing=0):
    return ("text", dict(x=x, y=y, w=w, h=h, text=value, size=size, bold=bold, color=color, align=align, spacing=spacing))


def heading(title, eyebrow, sub=None, w=12.1):
    """Eyebrow, title and optional subtitle, at the same place on every slide."""
    items = [text(0.6, 0.42, w, 0.3, eyebrow.upper(), 12, True, INDIGO, spacing=1.5),
             text(0.6, 0.74, w, 0.62, title, 32, True, INK)]
    if sub:
        items.append(text(0.6, 1.4, w, 0.35, sub, 15, color=SLATE))
    return items


def card(x, y, w, h, fill=WHITE, alpha=100, border=None, radius=0.14, shadow=True):
    return ("card", dict(x=x, y=y, w=w, h=h, fill=fill, alpha=alpha, border=border, radius=radius, shadow=shadow))


def badge(x, y, label, d=0.5, fill=INDIGO, color=WHITE, size=14):
    return [("circle", dict(x=x, y=y, d=d, fill=fill)), text(x, y + (d - size / 72 * 1.21) / 2, d, d, label, size, True, color, "center")]


def image(name, x, y, w=None, h=None, radius=0.08):
    return ("image", dict(path=str(ASSETS / name), x=x, y=y, w=w, h=h, radius=radius))


def rows(x, y, w, entries, pitch, lead_size=17, body_size=14, numbered=True):
    """Numbered (or dotted) rows: a badge, a bold lead line and an optional body line."""
    items = []
    for i, entry in enumerate(entries):
        lead, body = entry if isinstance(entry, tuple) else (entry, None)
        top = y + i * pitch
        items += badge(x, top, f"{i + 1:02d}" if numbered else "", 0.5 if numbered else 0.22, size=13)
        tx = x + (0.72 if numbered else 0.42)
        items.append(text(tx, top + (0.06 if numbered else -0.04), w - (tx - x), 0.5, lead, lead_size, True, INK))
        if body:
            items.append(text(tx, top + lead_size / 72 * 1.21 + 0.14, w - (tx - x), 0.5, body, body_size, color=SLATE))
    return items


def stat(x, y, w, h, value, label, dark=False, value_size=40, kicker=None):
    fill, value_color, label_color = (INDIGO, WHITE, "E0E7FF") if dark else (WHITE, INDIGO, SLATE)
    items = [card(x, y, w, h, fill)]
    top = y + 0.22
    if kicker:
        items.append(text(x + 0.3, top, w - 0.6, 0.3, kicker.upper(), 11, True, label_color, spacing=1.2))
        top += 0.34
    items.append(text(x + 0.3, top, w - 0.6, value_size / 72 * 1.25, value, value_size, True, value_color))
    items.append(text(x + 0.3, top + value_size / 72 * 1.21 + 0.1, w - 0.6, h, label, 14, color=label_color))
    return items


def note_card(x, y, w, h, title_, body, size=14, compact=False):
    pad, gap = (0.14, 0.42) if compact else (0.2, 0.6)
    return [card(x, y, w, h), text(x + 0.25, y + pad, w - 0.5, 0.35, title_, 15 if compact else 16, True, INK),
            text(x + 0.25, y + gap, w - 0.5, h - gap, body, size, color=SLATE)]


# Architecture diagram: nodes (id, label, sublabel, x, y, w, h, kind) and directed edges.
NODES = [
    ("clock", "Simulation clock", "persistent, pausable", 0.45, 1.95, 2.05, 0.72, "code"),
    ("sched", "Scheduler", "one-shot, restart-safe", 0.45, 3.15, 2.05, 0.72, "code"),
    ("planner", "Event planner", "AMD Token Factory", 3.1, 1.95, 2.05, 0.72, "model"),
    ("fallback", "Fallback planner", "deterministic, API-free", 3.1, 3.15, 2.05, 0.72, "code"),
    ("validator", "Validator", "schema + facts", 5.75, 2.55, 2.05, 0.72, "code"),
    ("state", "Canonical state", "persistent store, audit trail", 5.5, 3.85, 2.55, 0.95, "store"),
    ("override", "Central command", "override pending dispatch", 9.0, 2.55, 2.55, 0.72, "human"),
    ("ui", "Live Floor & Analytics", "UI, fixed snapshots", 9.0, 3.99, 2.55, 0.72, "code"),
    ("resolution", "Resolution engine", "hidden skill → outcome", 9.0, 5.4, 2.55, 0.72, "code"),
    ("context", "Context engine", "eligibility, travel, history", 5.75, 5.4, 2.05, 0.72, "code"),
    ("compass", "Compass", "ranks responders, fixes", 3.1, 5.05, 2.05, 0.72, "model"),
    ("baselines", "Baselines", "nearest + naive score", 3.1, 6.2, 2.05, 0.72, "code"),
    ("dispatch", "Dispatch", "ranked offer queue", 0.45, 5.6, 2.05, 0.72, "code"),
    ("responders", "Responders", "accept or reject", 0.45, 4.4, 2.05, 0.72, "human"),
]
EDGES = [
    ("clock", "sched"), ("sched", "planner"), ("sched", "fallback"),
    ("planner", "validator"), ("fallback", "validator"), ("validator", "state"),
    ("override", "state"), ("state", "ui"), ("resolution", "state"),
    ("state", "context"), ("context", "compass"), ("context", "baselines"),
    ("compass", "dispatch"), ("baselines", "dispatch"), ("dispatch", "responders"),
    ("responders", "state"),
]
# Autonomy loop: the system runs dispatch on its own; people decide at the amber nodes.
LOOP_NODES = [
    ("issue", "Issue raised", "equipment fault, symptom", 0.45, 2.0, 2.1, 0.72, "code"),
    ("rank", "Rank responders", "Compass on AMD", 3.05, 2.0, 2.1, 0.72, "model"),
    ("offer", "Offer sent", "automatic, ranked queue", 5.65, 2.0, 2.1, 0.72, "code"),
    ("accept", "Responder", "accepts or rejects", 8.25, 2.0, 2.1, 0.72, "human"),
    ("fix", "Rank fixes", "Compass on AMD", 10.85, 2.0, 2.1, 0.72, "model"),
    ("overtime", "Rank overtime", "Compass, legal options only", 0.45, 3.5, 2.1, 0.72, "model"),
    ("approve", "Central command", "approves or rejects", 3.05, 3.5, 2.1, 0.72, "human"),
    ("override", "Central command", "override, transfer, support", 8.25, 3.5, 2.1, 0.72, "human"),
    ("attempt", "Repair attempt", "outcome from hidden truth", 10.85, 3.5, 2.1, 0.72, "code"),
    ("backlog", "Backlog review", "before each handover", 0.45, 5.0, 2.1, 0.72, "code"),
    ("history", "Observed history", "evidence for every ranking", 8.25, 5.0, 2.1, 0.72, "store"),
    ("screens", "Wall screens", "floor, issues, feed, Pareto", 8.25, 6.25, 2.1, 0.72, "code"),
]
LOOP_EDGES = [
    ("issue", "rank"), ("rank", "offer"), ("offer", "accept"), ("accept", "fix"), ("fix", "attempt"),
    ("attempt", "history"), ("history", "rank"), ("override", "offer"), ("history", "screens"),
    ("history", "backlog"), ("backlog", "overtime"), ("overtime", "approve"), ("approve", "rank"),
]
NODE_STYLE = {  # fill, border, title colour, sub colour
    "code": (WHITE, "98A2B3", INK, SLATE),
    "model": (INDIGO, INDIGO, WHITE, "E0E7FF"),
    "store": (INDIGO_DEEP, INDIGO_DEEP, WHITE, "C7D2FE"),
    "human": (AMBER_WASH, AMBER, INK, SLATE),
}
LEGEND = [("Application code", "code", 1.45), ("AMD model", "model", 1.0), ("Store", "store", 0.55), ("People", "human", 0.65)]


def legend_entries(spec):
    """(label, kind, x, y) per entry, laid out left to right."""
    x, entries = spec["x"], []
    for label, kind, width in LEGEND:
        entries.append((label, kind, x, spec["y"]))
        x += 0.42 + width + 0.3
    return entries


def diagram(nodes, edges):
    """The legend runs along the top right, level with the eyebrow."""
    return [("diagram", dict(nodes=nodes, edges=edges)), ("legend", dict(x=6.2, y=0.47))]


def wall():
    """Six monitors, one panel each, labelled with a chip on the screen."""
    screens = [("screen-map.png", "Floor map"), ("screen-issues.png", "Issues"), ("screen-dispatch.png", "Dispatch"),
               ("screen-events.png", "Event feed"), ("screen-pareto.png", "Issue Pareto"), ("screen-analytics.png", "Analytics")]
    items, w, gap = [], 3.75, 0.425
    h = w * 9 / 16
    for i, (name, label) in enumerate(screens):
        x, y = 0.6 + (i % 3) * (w + gap), 2.0 + (i // 3) * (h + 0.32)
        items += [card(x - 0.08, y - 0.08, w + 0.16, h + 0.16, INDIGO_DEEP, radius=0.1),
                  image(name, x, y, w=w, h=h, radius=0.04),
                  card(x + w - 1.37, y + h - 0.42, 1.25, 0.3, INDIGO_DEEP, alpha=90, radius=0.15, shadow=False),
                  text(x + w - 1.37, y + h - 0.37, 1.25, 0.25, label, 11, True, WHITE, "center")]
    return items


def chevrons(y, steps):
    """A left-to-right process strip; a step marked human gets the amber people style."""
    items, w, overlap = [], 3.17, 0.18
    for i, (label, human) in enumerate(steps):
        x = 0.6 + i * (w - overlap)
        fill, color = (AMBER_WASH, INK) if human else ((INDIGO, WHITE) if i == 2 else (INDIGO_WASH, INK))
        items += [("chevron", dict(x=x, y=y, w=w, h=0.72, fill=fill, border=AMBER if human else None, first=i == 0)),
                  text(x + 0.4, y + 0.24, w - 0.75, 0.3, label, 15, True, color, "center")]
    return items


def weekly_bars(x, y, w, h):
    """September first-time fix by week, on a full 0-100% scale (no truncated axis)."""
    weeks = [("Week 1", 67.1), ("Week 2", 74.9), ("Week 3", 76.5), ("Week 4", 75.6)]
    items = [card(x, y, w, h), text(x + 0.3, y + 0.22, w - 0.6, 0.3, "First-time fix by week, September", 16, True, INK),
             text(x + 0.3, y + 0.56, w - 0.6, 0.3, "Rose, then dipped: not a steady climb. Bars on a full 0–100% track.", 13, color=SLATE)]
    base, top = y + h - 0.5, y + 1.0
    bar_w, pitch = 0.95, (w - 0.6) / len(weeks)
    for i, (label, value) in enumerate(weeks):
        bx = x + 0.3 + i * pitch + (pitch - bar_w) / 2
        bh = (base - top) * value / 100
        items += [card(bx, top, bar_w, base - top, INDIGO_WASH, radius=0.06, shadow=False),
                  card(bx, base - bh, bar_w, bh, INDIGO if i == 2 else "A5B4FC", radius=0.06, shadow=False),
                  text(bx - 0.2, base - bh + 0.1, bar_w + 0.4, 0.3, f"{value}%", 13, True, WHITE, "center"),
                  text(bx - 0.2, base + 0.1, bar_w + 0.4, 0.3, label, 12, color=SLATE, align="center")]
    return items


SLIDES = [
    dict(bg="cover", items=[],
         notes="ResponseCompass is a dispatch agent for operations control rooms, built on AMD Token Factory. It runs "
               "dispatch on its own and the control room supervises it: Compass ranks who to send and which fix to try "
               "first, from the plant's own repair history, and people approve or override the decisions that matter."),

    dict(bg="wallpaper", items=[
        *heading("Closest is not best", "The problem", w=6.0),
        *rows(0.6, 1.85, 5.9, [
            ("Nearest is not most experienced.", "The closest technician may never have fixed this fault."),
            ("The first fix is often a guess.", "A wrong guess means a reopened issue and another trip."),
            ("Backlogs force overtime calls.", "Who may legally stay on, and who covers the most work?"),
            ("Know-how is locked away.", "It lives in people's heads and old tickets, not in dispatch."),
        ], pitch=1.2)],
         notes="Proximity is easy to compute and easy to defend, but it ignores fit. What matters is who is likely to "
               "resolve this class of issue on the first attempt, with which procedure, and whether the shift can "
               "cover the backlog."),

    dict(items=[*heading("One agent, many screens", "Control room",
                         "Each panel runs full screen on its own monitor; the operator console keeps the controls."),
                *wall()],
         notes="A control room puts the floor map, the issue queue, the dispatch context, the event feed, the Issue "
               "Pareto and analytics on separate monitors (/rooms/<room>/display/<screen>). The screens update live from "
               "the same state. The operator console, with override and approval controls, is the seventh screen."),

    dict(items=[*heading("Autonomous by default, human in command", "How it runs",
                         "Code runs the loop; AMD models rank; people decide at the amber steps."),
                *diagram(LOOP_NODES, LOOP_EDGES)],
         notes="Issues are dispatched without anyone touching the console: Compass ranks the eligible responders, an "
               "offer goes out, the responder accepts, Compass ranks the troubleshooting procedures, and the repair "
               "outcome feeds the history that every later ranking uses. Central command can override, transfer or add "
               "support at any time. Before each shift handover a backlog review asks Compass to rank lawful overtime "
               "options; central command approves or rejects them, and nothing changes without that decision."),

    dict(items=[*heading("Code owns the facts. Models only rank.", "Architecture",
                         "Application code owns time, facts, eligibility, baselines and every mutation; models only propose."),
                *diagram(NODES, EDGES)],
         notes="The ranker receives application-owned evidence and ranks only the eligible candidates, fixes or "
               "overtime options it is given. The planner proposes synthetic incidents; code validates every proposal "
               "and owns all mutations. If a model is slow, rate-limited or returns invalid output, the deterministic "
               "fallback takes over, so dispatch continues."),

    dict(bg="wallpaper", items=[
        *heading("Honest by design", "Trust", w=6.0),
        *note_card(0.6, 1.75, 2.85, 2.3, "Hidden truth", "True skill and true fault cause stay hidden from every model, screen and chart."),
        *note_card(3.7, 1.75, 2.85, 2.3, "Outcomes from truth", "Repairs succeed or fail from that hidden truth, not from what a model claims."),
        *note_card(0.6, 4.3, 2.85, 2.3, "Baselines alongside", "Nearest, naive score and most common fix run next to Compass; disagreements are recorded."),
        *note_card(3.7, 4.3, 2.85, 2.3, "Fair analytics", "Denominators shown, thin samples suppressed, no leaderboard. Pay, seniority and HR data are never inputs."),
    ], notes="We wanted the system to be able to show Compass being wrong as well as right. Hidden truth plus baselines "
             "make its choices inspectable. Disagreement alone does not prove better outcomes."),

    dict(items=[*heading("The right responder", "Live capture",
                         "ISS-3618, Station 48 · ranked live on AMD Token Factory (MiMo-V2.6-Flash)"),
                *stat(0.6, 2.0, 3.6, 1.95, "26", "similar cases · Engr361, 0.5 min away", kicker="Nearest-responder baseline"),
                *stat(4.5, 2.0, 4.3, 1.95, "110", "similar cases, 87% success · Engr370, 0.5 min away", dark=True, kicker="Compass on AMD"),
                *note_card(9.1, 2.0, 3.6, 1.95, "4× the record", "Same half-minute trip. The offer went out automatically; central command could override it.", size=13),
                image("compass-record.png", 1.07, 4.25, w=11.2)],
         notes="Captured live on 2026-10-18 (simulated 6 October). The nearest-responder baseline chose Engr361; "
               "Compass ranked Engr370 first. One good outcome is not proof of better performance; disagreements with "
               "the baselines are recorded for review. In this capture 15 of 18 live responder rankings returned valid "
               "output; the rest fell back deterministically."),

    dict(items=[*heading("The right fix", "Live capture",
                         "Same record: Compass ranks the troubleshooting procedures for the observed symptom"),
                *stat(0.6, 2.0, 4.4, 1.6, "7.0 min", "fixed first time with the top-ranked procedure", dark=True),
                *rows(0.6, 3.9, 4.4, [
                    ("Evidence from code", "Procedures and observed success."),
                    ("Compass ranks", "Mechanical alignment came first."),
                    ("Fail, then next", "A failed attempt reopens the issue."),
                ], pitch=0.95, lead_size=16),
                image("fix-record.png", 5.35, 2.0, w=7.35),
                *note_card(5.35, 5.6, 7.35, 1.05, "Agrees with the baseline",
                           "The most common fix for this symptom also came first; disagreements are recorded.", size=13, compact=True)],
         notes="The true fault cause is hidden simulation truth, like responder skill: the model sees only the public "
               "symptom and observed outcomes. Compass agreed with the most-common-fix baseline here. In this capture "
               "12 of 15 live fix rankings were valid AMD results. A live check also found and fixed a defect: the fix "
               "ranking never reached the provider until that check."),

    dict(items=[*heading("Overtime, decided by people", "Staffing",
                         "Nothing changes until central command approves; unanswered recommendations expire."),
                *chevrons(1.98, [("Backlog review", False), ("Lawful options only", False), ("Compass ranks", False), ("Command approves", True)]),
                *rows(0.6, 3.2, 5.9, [
                    ("Legal limits first", "16 h duty, 12 h rest and the staffing cap, checked in code."),
                    ("Ranked by coverage", "Observed coverage of the backlog; first-eligible is the baseline."),
                    ("Audited decisions", "Every approval or rejection, and who made it, is recorded."),
                ], pitch=1.15, lead_size=16),
                image("staffing.png", 7.05, 3.05, h=3.6)],
         notes="Live capture: staffing-7 was ranked by MiMo on AMD Token Factory and approved by central command in the "
               "recording. A live check found that approving went stale whenever the top option had picked up a job in "
               "the meantime; extensions now depend only on duty and rest limits. Rooms can opt into an audited "
               "auto-approve policy; the default is operator approval."),

    dict(items=[*heading("Where the downtime comes from", "Issue Pareto",
                         "Issue class → observed symptom → applied fix, or floor → station"),
                image("screen-pareto.png", 0.6, 2.0, w=8.2),
                *[item for i, (t, b) in enumerate([
                    ("Three measures", "Issues, downtime or failed attempts, by day, week or month."),
                    ("The vital few", "Sorted bars, cumulative line and an 80% marker."),
                    ("Drill down", "Class to symptom to the fix that was applied."),
                    ("Public records only", "Hidden causes are never used."),
                ]) for item in note_card(9.1, 2.0 + i * 1.18, 3.6, 1.04, t, b, size=13, compact=True)]],
         notes="The Pareto answers the supervisor's first question: which few problems cost us most of the time. In "
               "September two of three issue classes accounted for 80% of issues; drilling into a symptom shows which "
               "procedures failed most often. It describes recorded work; it does not rank people or claim causation."),

    dict(items=[*heading("A month of operations", "Results", "e8-aura-p1, simulated 1 September to 6 October"),
                *stat(0.6, 1.95, 3.83, 1.6, "3,512", "issues in the month-long history"),
                *stat(4.735, 1.95, 3.83, 1.6, "73.5%", "first-time fix in September (1,858 of 2,527)", dark=True),
                *stat(8.87, 1.95, 3.83, 1.6, "2,966 / 2,985", "issues raised / resolved in September", value_size=34),
                *weekly_bars(0.6, 3.8, 7.85, 2.9),
                *note_card(8.75, 3.8, 3.95, 2.9, "How it was made",
                           "Generated with the app's own engine, offline, with deterministic fallbacks. Every record is one the app "
                           "would produce; it is not an AMD-only benchmark. Weeks start Monday 08:00 UTC.", size=14)],
         notes="This history was generated with the app's own engine, offline and deterministically, so every record "
               "is one the app itself would produce. The month extension used deterministic fallbacks, so it is not an "
               "AMD-only benchmark or a controlled comparison of model and baseline outcomes. Weekly rates use Monday "
               "08:00 UTC boundaries; they rose then dipped, and we report them as measured."),

    dict(items=[*heading("Built on AMD Token Factory", "Inference",
                         "MiMo-V2.6-Flash ranks responders, fixes and overtime; the event planner can use another model."),
                *stat(0.6, 1.95, 3.83, 1.6, "9 / 10", "valid responder rankings in the recorded MiMo check", dark=True),
                *stat(4.735, 1.95, 3.83, 1.6, "10 / 10", "valid event plans in the same check"),
                *stat(8.87, 1.95, 3.83, 1.6, "10 / min", "local request budget, 2 concurrent"),
                *rows(0.6, 3.95, 5.6, [
                    ("Validators are the contract", "OpenAI-compatible JSON, checked by code."),
                    ("Bounded failure", "Two retries on 429 only, then a circuit breaker."),
                    ("Never holds the clock", "Model calls run in a background lane."),
                ], pitch=0.95, lead_size=16),
                card(6.55, 3.95, 6.15, 2.45),
                text(6.8, 4.15, 5.7, 0.3, "Each role picks its own model", 16, True, INK),
                image("settings-models.png", 6.8, 4.6, w=5.65, radius=0.05)],
         notes="We measured several Token Factory models on the same scenario and chose the default per role from the "
               "data. These are small-sample validity results, not outcome accuracy or uptime guarantees. Retries: at "
               "most two, only on HTTP 429. Five consecutive service failures open the breaker; one probe after 60 s."),

    dict(bg="wallpaper", items=[
        *heading("Engineering quality", "Built to run", w=6.0),
        *stat(0.6, 1.75, 2.85, 1.75, "334", "hermetic unit and integration tests", value_size=44),
        *stat(3.7, 1.75, 2.85, 1.75, "72", "Chromium end-to-end tests", dark=True, value_size=44),
        *rows(0.6, 3.85, 5.95, [
            ("One app, two modes", "Offline makes no API calls and is fully reproducible."),
            ("Crash-safe state", "Checksummed journal and a restart-safe scheduler."),
            ("Spend only when watched", "An unattended room pauses itself after 10 idle minutes."),
        ], pitch=0.95, lead_size=16)],
         notes="Everything that matters is tested without spending API credits; the browser tests force provider "
               "failures to prove the fallbacks. A live check on AMD found and fixed one defect the hermetic tests "
               "could not: the fix ranking never reached the provider. Runs on Node.js 22 or Docker, with no services "
               "besides the model API."),

    dict(bg="wallpaper", items=[
        *heading("What's next", "Roadmap", w=6.0),
        *rows(0.6, 1.8, 5.9, [
            ("Real systems", "Maintenance, ticketing and rostering in place of the simulated floor."),
            ("Learn from approvals", "Feed operator decisions and outcomes back into the evidence."),
            ("Longer live measurement", "Compass against its baselines on live AMD rankings."),
        ], pitch=1.05, lead_size=16),
        text(0.6, 5.15, 6.0, 0.5, "The right responder. The right fix.", 26, True, INDIGO),
        text(0.6, 5.72, 6.0, 0.4, "Runs autonomously. Supervised by your control room.", 16, True, INK),
        text(0.6, 6.1, 6.0, 0.4, "Built on AMD Token Factory.", 16, color=SLATE),
    ], notes="Responder ranking, fix ranking, overtime recommendations and the control-room screens are built. Next is "
             "real integration and a longer live measurement."),
]

# ---------------------------------------------------------------------------
# Helpers shared by both renderers
# ---------------------------------------------------------------------------

def runs_of(value):
    """Splits **bold** markup into (text, bold) runs."""
    parts = value.split("**")
    return [(part, i % 2 == 1) for i, part in enumerate(parts) if part]


def node_map(nodes):
    return {n[0]: dict(id=n[0], label=n[1], sub=n[2], x=n[3], y=n[4], w=n[5], h=n[6], kind=n[7]) for n in nodes}


def sites(node):
    """Connection sites of a rectangle in PowerPoint order: top, left, bottom, right."""
    x, y, w, h = node["x"], node["y"], node["w"], node["h"]
    return [(x + w / 2, y), (x, y + h / 2), (x + w / 2, y + h), (x + w, y + h / 2)]


def pick_sites(a, b):
    ax, ay = a["x"] + a["w"] / 2, a["y"] + a["h"] / 2
    bx, by = b["x"] + b["w"] / 2, b["y"] + b["h"] / 2
    dx, dy = bx - ax, by - ay
    if abs(dx) >= abs(dy) * 1.2:
        return (3, 1) if dx > 0 else (1, 3)
    return (2, 0) if dy > 0 else (0, 2)


def image_size(path, w, h):
    with Image.open(path) as img:
        iw, ih = img.size
    if w and not h:
        h = w * ih / iw
    elif h and not w:
        w = h * iw / ih
    return w, h


def background_path(spec):
    return {"cover": COVER, "wallpaper": WALLPAPER}.get(spec.get("bg"), SOFT_WALLPAPER)


def build_soft_wallpaper():
    """The wallpaper washed towards the page colour, with its wordmark kept at full strength."""
    CACHE.mkdir(parents=True, exist_ok=True)
    with Image.open(WALLPAPER) as source:
        original = source.convert("RGB")
    soft = Image.blend(original, Image.new("RGB", original.size, "#" + BG), 0.7)
    mark = (0, int(original.height * 0.895), int(original.width * 0.2), int(original.height * 0.965))
    soft.paste(original.crop(mark), mark[:2])
    soft.save(SOFT_WALLPAPER)

# ---------------------------------------------------------------------------
# PowerPoint renderer
# ---------------------------------------------------------------------------

def rgb(hex_):
    return RGBColor.from_string(hex_)


def set_font(run, size, bold=False, color=INK, spacing=0):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    if spacing:
        run._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))


def soft_shadow(shape_element):
    sp_pr = shape_element.find(qn("p:spPr"))
    # shadow.inherit = False has already written an empty effectLst; a second one is invalid.
    effects = sp_pr.find(qn("a:effectLst"))
    if effects is None:
        effects = etree.SubElement(sp_pr, qn("a:effectLst"))
    shadow = etree.SubElement(effects, qn("a:outerShdw"), blurRad="203200", dist="38100", dir="5400000", algn="t", rotWithShape="0")
    color = etree.SubElement(shadow, qn("a:srgbClr"), val=INDIGO_DEEP)
    etree.SubElement(color, qn("a:alpha"), val="13000")


def pptx_text(slide, spec):
    box = slide.shapes.add_textbox(Inches(spec["x"]), Inches(spec["y"]), Inches(spec["w"]), Inches(spec["h"]))
    frame = box.text_frame
    frame.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, side, 0)
    para = frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER if spec.get("align") == "center" else PP_ALIGN.LEFT
    for value, bold in runs_of(spec["text"]):
        run = para.add_run()
        run.text = value
        set_font(run, spec["size"], bold or spec.get("bold", False), spec.get("color", INK), spec.get("spacing", 0))


def pptx_card(slide, spec):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(spec["x"]), Inches(spec["y"]), Inches(spec["w"]), Inches(spec["h"]))
    shape.adjustments[0] = min(0.5, spec["radius"] / min(spec["w"], spec["h"]))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(spec["fill"])
    if spec["alpha"] < 100:
        fill_color = shape._element.find(qn("p:spPr")).find(qn("a:solidFill"))[0]
        etree.SubElement(fill_color, qn("a:alpha"), val=str(spec["alpha"] * 1000))
    if spec["border"]:
        shape.line.color.rgb = rgb(spec["border"])
        shape.line.width = Pt(1.25)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    if spec["shadow"]:
        soft_shadow(shape._element)


def pptx_circle(slide, spec):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(spec["x"]), Inches(spec["y"]), Inches(spec["d"]), Inches(spec["d"]))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(spec["fill"])
    shape.line.fill.background()
    shape.shadow.inherit = False


def pptx_chevron(slide, spec):
    kind = MSO_SHAPE.PENTAGON if spec["first"] else MSO_SHAPE.CHEVRON
    shape = slide.shapes.add_shape(kind, Inches(spec["x"]), Inches(spec["y"]), Inches(spec["w"]), Inches(spec["h"]))
    shape.adjustments[0] = 0.3 * spec["h"] / min(spec["w"], spec["h"])
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(spec["fill"])
    if spec["border"]:
        shape.line.color.rgb = rgb(spec["border"])
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False


def pptx_image(slide, spec):
    w, h = image_size(spec["path"], spec["w"], spec["h"])
    pic = slide.shapes.add_picture(spec["path"], Inches(spec["x"]), Inches(spec["y"]), Inches(w), Inches(h))
    geometry = pic._element.find(qn("p:spPr")).find(qn("a:prstGeom"))
    geometry.set("prst", "roundRect")
    av = geometry.find(qn("a:avLst"))
    etree.SubElement(av, qn("a:gd"), name="adj", fmla=f"val {int(spec['radius'] / min(w, h) * 100000)}")
    pic.line.color.rgb = rgb(LINE)
    pic.line.width = Pt(0.75)
    soft_shadow(pic._element)


def add_arrowhead(connector):
    ln = connector.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")


def pptx_diagram(slide, spec):
    shapes = {}
    nodes = node_map(spec["nodes"])
    for node in nodes.values():
        fill, border, title_color, sub_color = NODE_STYLE[node["kind"]]
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(node["x"]), Inches(node["y"]), Inches(node["w"]), Inches(node["h"]))
        shape.adjustments[0] = 0.18
        shape.name = f"Node: {node['label']}"
        shape.shadow.inherit = False
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.color.rgb = rgb(border)
        shape.line.width = Pt(1.5)
        soft_shadow(shape._element)
        frame = shape.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for side in ("margin_left", "margin_right"):
            setattr(frame, side, Inches(0.06))
        frame.margin_top = frame.margin_bottom = Inches(0.02)
        head = frame.paragraphs[0]
        head.alignment = PP_ALIGN.CENTER
        run = head.add_run()
        run.text = node["label"]
        set_font(run, 13, True, title_color)
        sub = frame.add_paragraph()
        sub.alignment = PP_ALIGN.CENTER
        run = sub.add_run()
        run.text = node["sub"]
        set_font(run, 9.5, False, sub_color)
        shapes[node["id"]] = shape
    for a, b in spec["edges"]:
        start, end = pick_sites(nodes[a], nodes[b])
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
        connector.begin_connect(shapes[a], start)
        connector.end_connect(shapes[b], end)
        connector.name = f"Arrow: {nodes[a]['label']} to {nodes[b]['label']}"
        connector.line.color.rgb = rgb(MUTED)
        connector.line.width = Pt(1.5)
        add_arrowhead(connector)


def pptx_legend(slide, spec):
    for label, kind, x, y in legend_entries(spec):
        fill, border, _, _ = NODE_STYLE[kind]
        swatch = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(0.32), Inches(0.2))
        swatch.fill.solid()
        swatch.fill.fore_color.rgb = rgb(fill)
        swatch.line.color.rgb = rgb(border)
        swatch.shadow.inherit = False
        pptx_text(slide, dict(x=x + 0.42, y=y - 0.01, w=1.8, h=0.28, text=label, size=12, color=SLATE))


def build_pptx():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    render = {"text": pptx_text, "card": pptx_card, "circle": pptx_circle, "chevron": pptx_chevron,
              "image": pptx_image, "diagram": pptx_diagram, "legend": pptx_legend}
    for number, spec in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        background = slide.shapes.add_picture(str(background_path(spec)), 0, 0, Emu(prs.slide_width), Emu(prs.slide_height))
        background.name = "Background"
        for kind, item in spec["items"]:
            render[kind](slide, item)
        if spec.get("bg") != "cover":
            pptx_text(slide, dict(x=12.2, y=7.0, w=0.5, h=0.3, text=str(number), size=11, color=MUTED, align="center"))
        slide.notes_slide.notes_text_frame.text = spec["notes"]
    prs.save(OUT_PPTX)

# ---------------------------------------------------------------------------
# PDF renderer: the same geometry as HTML, printed by Chromium
# ---------------------------------------------------------------------------

def css_color(hex_, alpha=100):
    r, g, b = (int(hex_[i:i + 2], 16) for i in (0, 2, 4))
    return f"rgba({r},{g},{b},{alpha / 100})"


def box(x, y, w, h):
    return f"left:{x}in;top:{y}in;width:{w}in;height:{h}in;"


SHADOW = "box-shadow:0 0.04in 0.22in rgba(30,27,75,0.13);"


def html_text(spec):
    runs = "".join(f"<b>{html.escape(t)}</b>" if b else html.escape(t) for t, b in runs_of(spec["text"]))
    weight = 700 if spec.get("bold") else 400
    return (f'<div class="t" style="{box(spec["x"], spec["y"], spec["w"], spec["h"])}font-size:{spec["size"]}pt;'
            f'font-weight:{weight};color:#{spec.get("color", INK)};text-align:{spec.get("align", "left")};'
            f'letter-spacing:{spec.get("spacing", 0)}pt">{runs}</div>')


def html_card(spec):
    border = f"border:1.25pt solid #{spec['border']};" if spec["border"] else ""
    return (f'<div style="{box(spec["x"], spec["y"], spec["w"], spec["h"])}background:{css_color(spec["fill"], spec["alpha"])};'
            f'border-radius:{spec["radius"]}in;{border}{SHADOW if spec["shadow"] else ""}"></div>')


def html_circle(spec):
    return f'<div style="{box(spec["x"], spec["y"], spec["d"], spec["d"])}background:#{spec["fill"]};border-radius:50%"></div>'


def html_chevron(spec):
    tip = 0.3 * spec["h"]
    tip_pct = tip / spec["w"] * 100
    points = ("0 0" if spec["first"] else "0 0") + f",{100 - tip_pct}% 0,100% 50%,{100 - tip_pct}% 100%,0 100%" + \
        ("" if spec["first"] else f",{tip_pct}% 50%")
    border = spec["border"]
    outer = (f'<div style="{box(spec["x"], spec["y"], spec["w"], spec["h"])}clip-path:polygon({points});'
             f'background:#{border or spec["fill"]}"></div>')
    if not border:
        return outer
    inset = 0.025
    return outer + (f'<div style="{box(spec["x"] + inset * 2, spec["y"] + inset, spec["w"] - inset * 3, spec["h"] - inset * 2)}'
                    f'clip-path:polygon({points});background:#{spec["fill"]}"></div>')


def html_image(spec):
    w, h = image_size(spec["path"], spec["w"], spec["h"])
    return (f'<img src="{Path(spec["path"]).as_uri()}" style="{box(spec["x"], spec["y"], w, h)}'
            f'border-radius:{spec["radius"]}in;outline:0.75pt solid #{LINE};outline-offset:-0.75pt;{SHADOW}">')


def html_diagram(spec):
    nodes = node_map(spec["nodes"])
    lines = []
    for a, b in spec["edges"]:
        s, e = pick_sites(nodes[a], nodes[b])
        (x1, y1), (x2, y2) = sites(nodes[a])[s], sites(nodes[b])[e]
        lines.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="#{MUTED}" stroke-width="0.021" marker-end="url(#arrow)"/>')
    svg = (f'<svg style="position:absolute;left:0;top:0;width:{W}in;height:{H}in" viewBox="0 0 {W} {H}">'
           f'<defs><marker id="arrow" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="5" markerHeight="5" orient="auto">'
           f'<path d="M0,0 L10,5 L0,10 z" fill="#{MUTED}"/></marker></defs>{"".join(lines)}</svg>')
    parts = [svg]
    for node in nodes.values():
        fill, border, title_color, sub_color = NODE_STYLE[node["kind"]]
        parts.append(
            f'<div class="node" style="{box(node["x"], node["y"], node["w"], node["h"])}background:#{fill};'
            f'border:1.5pt solid #{border};{SHADOW}"><div style="font-size:13pt;font-weight:700;color:#{title_color}">'
            f'{html.escape(node["label"])}</div><div style="font-size:9.5pt;color:#{sub_color}">{html.escape(node["sub"])}</div></div>')
    return "".join(parts)


def html_legend(spec):
    parts = []
    for label, kind, x, y in legend_entries(spec):
        fill, border, _, _ = NODE_STYLE[kind]
        parts.append(f'<div style="{box(x, y, 0.32, 0.2)}background:#{fill};border:1pt solid #{border};border-radius:0.05in"></div>')
        parts.append(html_text(dict(x=x + 0.42, y=y - 0.01, w=1.8, h=0.28, text=label, size=12, color=SLATE)))
    return "".join(parts)


def build_pdf():
    render = {"text": html_text, "card": html_card, "circle": html_circle, "chevron": html_chevron,
              "image": html_image, "diagram": html_diagram, "legend": html_legend}
    pages = []
    for number, spec in enumerate(SLIDES, start=1):
        body = "".join(render[kind](item) for kind, item in spec["items"])
        if spec.get("bg") != "cover":
            body += html_text(dict(x=12.2, y=7.0, w=0.5, h=0.3, text=str(number), size=11, color=MUTED, align="center"))
        pages.append(f'<section style="background-image:url({background_path(spec).as_uri()})">{body}</section>')
    document = f"""<!doctype html><html><head><meta charset="utf-8"><style>
@font-face {{ font-family: Inter; src: url({INTER.as_uri()}) format("woff2"); font-weight: 100 900; }}
@page {{ size: {W}in {H}in; margin: 0; }}
* {{ box-sizing: border-box; margin: 0; }}
body {{ font-family: Inter, sans-serif; -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
section {{ position: relative; width: {W}in; height: {H}in; overflow: hidden; background-size: 100% 100%; break-after: page; }}
section > * {{ position: absolute; }}
.t {{ line-height: 1.21; overflow-wrap: break-word; }}
.t b {{ font-weight: 700; }}
.node {{ border-radius: 0.13in; display: flex; flex-direction: column; align-items: center; justify-content: center; text-align: center; line-height: 1.2; }}
</style></head><body>{"".join(pages)}</body></html>"""
    html_path = CACHE / "deck.html"
    html_path.write_text(document, encoding="utf-8")
    subprocess.run(["node", str(Path(__file__).with_name("print-pdf.mjs")), str(html_path), str(OUT_PDF)], check=True, cwd=REPO)


if __name__ == "__main__":
    build_soft_wallpaper()
    build_pptx()
    build_pdf()
    print(f"wrote {OUT_PPTX.relative_to(REPO)} and {OUT_PDF.relative_to(REPO)}")
